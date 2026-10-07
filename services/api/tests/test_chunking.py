import io
from itertools import pairwise

from docx import Document as DocxDocument

from app.modules.knowledge.chunking import chunk_sections, count_tokens
from app.modules.knowledge.parsing import DOCX, MD, Section, parse_document

LONG_TEXT = " ".join(f"word{i}" for i in range(3000))


def test_chunks_respect_size_and_overlap() -> None:
    chunks = chunk_sections([Section(text=LONG_TEXT)], chunk_tokens=500, overlap_tokens=60)
    assert len(chunks) > 3
    for chunk in chunks:
        assert 0 < chunk.token_count <= 500
        assert chunk.token_count == count_tokens(chunk.content)

    for prev, nxt in pairwise(chunks):
        prev_words, next_words = prev.content.split(), nxt.content.split()
        # The next chunk starts with the tail of the previous one.
        overlap = next(
            n for n in range(len(next_words), 0, -1) if prev_words[-n:] == next_words[:n]
        )
        overlap_tokens = count_tokens(" ".join(next_words[:overlap]))
        assert 30 <= overlap_tokens <= 65


def test_chunks_never_split_words_and_cover_everything() -> None:
    chunks = chunk_sections([Section(text=LONG_TEXT)], chunk_tokens=200, overlap_tokens=30)
    original = set(LONG_TEXT.split())
    seen: set[str] = set()
    for chunk in chunks:
        words = chunk.content.split()
        assert set(words) <= original
        seen.update(words)
    assert seen == original


def test_heading_path_from_markdown() -> None:
    md = (
        "# Returns\n\nIntro to returns.\n\n"
        "## Refund window\n\nRefunds within 30 days.\n\n"
        "### Exceptions\n\nSale items are final.\n\n"
        "## Shipping back\n\nUse the prepaid label.\n\n"
        "```\n# not a heading inside code\n```\n"
    )
    chunks = chunk_sections(parse_document(md.encode(), MD))
    paths = [(c.metadata["heading_path"], c.content.split("\n")[0]) for c in chunks]
    assert paths[:4] == [
        (["Returns"], "Intro to returns."),
        (["Returns", "Refund window"], "Refunds within 30 days."),
        (["Returns", "Refund window", "Exceptions"], "Sale items are final."),
        (["Returns", "Shipping back"], "Use the prepaid label."),
    ]
    assert "# not a heading inside code" in chunks[3].content
    assert all(c.metadata["page_number"] is None for c in chunks)


def test_no_empty_chunks() -> None:
    sections = [
        Section(text=""),
        Section(text="   \n\n  "),
        Section(text="Real content.", heading_path=["A"]),
    ]
    chunks = chunk_sections(sections)
    assert [c.content for c in chunks] == ["Real content."]

    # Headings with no body produce no chunks.
    md = "# Empty\n\n## Also empty\n\n## Has text\n\nHello there.\n"
    chunks = chunk_sections(parse_document(md.encode(), MD))
    assert [(c.metadata["heading_path"], c.content) for c in chunks] == [
        (["Empty", "Has text"], "Hello there.")
    ]


def test_docx_heading_styles() -> None:
    doc = DocxDocument()
    doc.add_heading("Policies", level=1)
    doc.add_paragraph("General policy text.")
    doc.add_heading("Warranty", level=2)
    doc.add_paragraph("Two year warranty.")
    buf = io.BytesIO()
    doc.save(buf)

    sections = parse_document(buf.getvalue(), DOCX)
    assert [(s.heading_path, s.text) for s in sections] == [
        (["Policies"], "General policy text."),
        (["Policies", "Warranty"], "Two year warranty."),
    ]
